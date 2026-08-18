from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SAGE_BLUE = "2B579A"
SAGE_HEADER = "5B6770"
INPUT_FILL = PatternFill("solid", fgColor="FFFFFF")
LABEL_FILL = PatternFill("solid", fgColor="E7E6E6")
TITLE_FILL = PatternFill("solid", fgColor=SAGE_BLUE)
GRID_HEADER = PatternFill("solid", fgColor=SAGE_HEADER)
THIN = Side(style="thin", color="A0A0A0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_simulator_config(root: Path) -> dict[str, Any]:
    path = root / "config" / "sage_simulator.json"
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def load_invoice_sample(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def _format_date(value: str) -> str:
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%b %d, %Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%b %d, %Y")
        except ValueError:
            continue
    return value


def _style_input(ws, cell: str) -> None:
    ws[cell].fill = INPUT_FILL
    ws[cell].border = BOX
    ws[cell].alignment = Alignment(vertical="center")


def _style_label(ws, cell: str) -> None:
    ws[cell].font = Font(name="Segoe UI", size=10, bold=True)
    ws[cell].alignment = Alignment(horizontal="right", vertical="center")


def create_sage_template(root: Path, config: dict[str, Any] | None = None) -> Path:
    config = config or load_simulator_config(root)
    layout = config["excel_layout"]
    template_path = root / config["paths"]["template"]
    template_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = layout["sheet_name"]
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 28
    ws["A1"] = "Sales/Invoicing"
    ws["A1"].font = Font(name="Segoe UI", bold=True, color="FFFFFF", size=12)
    ws["A1"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    fields = [
        ("A3", "Customer ID", "B3", "H3"),
        ("A4", "Ship to", "B4", "H4"),
        ("A5", "Invoice date", "B5", "C5"),
        ("E5", "Due date", "F5", "H5"),
        ("A6", "Invoice No.", "B6", "D6"),
        ("A7", "Ship via", "B7", "C7"),
        ("E7", "A/R account", "F7", "H7"),
        ("A8", "Sales rep", "B8", "H8"),
    ]
    for label_cell, label_text, start, end in fields:
        ws[label_cell] = label_text
        _style_label(ws, label_cell)
        ws.merge_cells(f"{start}:{end}")
        _style_input(ws, start)

    headers = ["Quantity", "Item", "Description", "G/L Account", "Unit Price", "Tax", "Amount", "Job"]
    header_row = layout["line_header_row"]
    ws.row_dimensions[header_row].height = 22
    for idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=title)
        cell.font = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
        cell.fill = GRID_HEADER
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in range(layout["line_start_row"], layout["line_start_row"] + 10):
        ws.row_dimensions[r].height = 20
        fill = PatternFill("solid", fgColor="FFFFFF" if r % 2 else "F5F5F5")
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.fill = fill
            cell.font = Font(name="Segoe UI", size=10)

    ws["A22"] = "Balance:"
    ws["A23"] = "Credit limit:"
    ws["A24"] = "Credit status:"
    for addr in ("A22", "A23", "A24"):
        _style_label(ws, addr)

    ws["F24"] = "Sales tax (ITBMS):"
    ws["F25"] = "Invoice total:"
    ws["F26"] = "Net due:"
    for addr in ("F24", "F25", "F26"):
        _style_label(ws, addr)
        _style_input(ws, addr.replace("F", "G"))

    widths = {"A": 11, "B": 13, "C": 38, "D": 11, "E": 11, "F": 7, "G": 11, "H": 8}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(template_path)
    return template_path


def fill_sage_workbook(
    template_path: Path,
    invoice: dict[str, Any],
    config: dict[str, Any],
    output_path: Path,
    clear_first: bool = True,
) -> Path:
    layout = config["excel_layout"]
    wb = load_workbook(template_path)
    ws = wb[layout["sheet_name"]]

    if clear_first:
        for key, cell in layout["header_cells"].items():
            ws[cell] = ""
        if "ship_to" in layout["header_cells"]:
            ws[layout["header_cells"].get("ship_to", "B4")] = ""
        cols = layout["line_columns"]
        for r in range(layout["line_start_row"], layout["line_start_row"] + 10):
            for col in cols.values():
                ws[f"{col}{r}"] = ""
        for footer_key in ("sales_tax", "invoice_total", "net_due"):
            ws[layout["footer_cells"][footer_key]] = ""

    header_cells = layout["header_cells"]
    ws[header_cells["customer_id"]] = invoice.get("customer_id", "")
    if "ship_to" in header_cells:
        ws[header_cells["ship_to"]] = invoice.get("customer_id", "")
    ws[header_cells["invoice_date"]] = _format_date(str(invoice.get("invoice_date", "")))
    ws[header_cells["due_date"]] = _format_date(str(invoice.get("due_date", invoice.get("invoice_date", ""))))
    ws[header_cells["invoice_no"]] = invoice.get("invoice_no", "")
    ws[header_cells["ship_via"]] = invoice.get("ship_via", config["defaults"].get("ship_via", ""))
    ws[header_cells["ar_account"]] = invoice.get("ar_account", config["defaults"].get("ar_account", ""))
    if invoice.get("sales_rep"):
        ws[header_cells["sales_rep"]] = invoice.get("sales_rep", "")

    cols = layout["line_columns"]
    start_row = layout["line_start_row"]
    for offset, line in enumerate(invoice.get("lines", [])):
        row = start_row + offset
        ws[f"{cols['quantity']}{row}"] = line.get("quantity", "")
        ws[f"{cols['item']}{row}"] = line.get("item", "")
        ws[f"{cols['description']}{row}"] = line.get("description", "")
        ws[f"{cols['gl_account']}{row}"] = line.get("gl_account", config["defaults"].get("gl_account", ""))
        ws[f"{cols['unit_price']}{row}"] = line.get("unit_price", "")
        ws[f"{cols['tax']}{row}"] = line.get("tax", "")
        ws[f"{cols['amount']}{row}"] = line.get("amount", "")
        ws[f"{cols['job']}{row}"] = line.get("job", "")

    footer = layout["footer_cells"]
    ws[footer["sales_tax"]] = invoice.get("sales_tax", "")
    ws[footer["invoice_total"]] = invoice.get("invoice_total", "")
    ws[footer["net_due"]] = invoice.get("net_due", invoice.get("invoice_total", ""))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def dataframe_to_invoice(frame, config: dict[str, Any], invoice_number: str | None = None) -> dict[str, Any]:
    if frame.empty:
        raise ValueError("No hay filas para convertir a factura.")

    if invoice_number and "Invoice Number" in frame.columns:
        subset = frame[frame["Invoice Number"].astype(str) == str(invoice_number)]
        if not subset.empty:
            frame = subset

    if "Invoice Number" in frame.columns:
        unique = frame["Invoice Number"].astype(str).unique()
        if len(unique) > 1:
            raise ValueError(
                "Hay varias facturas en los datos. Selecciona una factura antes de usar Carga Sage."
            )

    first = frame.iloc[0]
    tax_map = config.get("tax_mapping", {})
    defaults = config.get("defaults", {})

    lines = []
    sales_tax = 0.0
    for _, row in frame.iterrows():
        tax_raw = str(row.get("Tax Code", defaults.get("codigo_impuesto", "1")))
        tax_code = tax_map.get(tax_raw, tax_raw)
        amount = float(row.get("Line Amount", 0) or 0)
        tax_amount = float(row.get("Tax Amount", 0) or 0)
        sales_tax += tax_amount
        item_val = row.get("Item", "")
        if not item_val or (isinstance(item_val, float) and pd.isna(item_val)):
            item_val = str(row.get("Description", ""))[:12]
        lines.append(
            {
                "quantity": float(row.get("Quantity", 0) or 0),
                "item": str(item_val),
                "description": str(row.get("Description", "")),
                "gl_account": str(row.get("GL Account", defaults.get("gl_account", "4001"))),
                "unit_price": float(row.get("Unit Price", 0) or 0),
                "tax": tax_code,
                "amount": amount,
                "job": "",
            }
        )

    invoice_total = float(first.get("Invoice Total", 0) or sum(l["amount"] for l in lines) or 0)
    return {
        "customer_id": str(first.get("Customer ID", first.get("Customer Name", ""))),
        "invoice_date": str(first.get("Date", "")),
        "due_date": str(first.get("Date", "")),
        "invoice_no": str(first.get("Invoice Number", "")),
        "ship_via": defaults.get("ship_via", "Airborne"),
        "ar_account": defaults.get("ar_account", "1101"),
        "sales_rep": "",
        "lines": lines,
        "sales_tax": round(sales_tax, 2) if sales_tax else round(invoice_total * 0.07, 2),
        "invoice_total": invoice_total,
        "net_due": invoice_total,
    }


def build_field_sequence(invoice: dict[str, Any], config: dict[str, Any]) -> list[dict[str, str]]:
    layout = config["excel_layout"]
    labels = config.get("field_labels", {})
    sequence: list[dict[str, str]] = []

    def add(cell: str, value: str, section: str, label: str | None = None) -> None:
        if value == "":
            return
        sequence.append(
            {
                "cell": cell,
                "value": value,
                "label": label or labels.get(cell, cell),
                "section": section,
            }
        )

    header_sources = {
        "customer_id": "customer_id",
        "ship_to": "customer_id",
        "invoice_date": "invoice_date",
        "due_date": "due_date",
        "invoice_no": "invoice_no",
        "ship_via": "ship_via",
        "ar_account": "ar_account",
        "sales_rep": "sales_rep",
    }
    for key, cell in layout["header_cells"].items():
        if key == "sales_rep" and not invoice.get("sales_rep"):
            continue
        source_key = header_sources.get(key, key)
        value = invoice.get(source_key, "")
        if key in ("invoice_date", "due_date"):
            value = _format_date(str(value))
        add(cell, str(value), "header")

    cols = layout["line_columns"]
    start_row = layout["line_start_row"]
    col_labels = {
        "quantity": "Quantity",
        "item": "Item",
        "description": "Description",
        "gl_account": "G/L Account",
        "unit_price": "Unit Price",
        "tax": "Tax",
        "amount": "Amount",
        "job": "Job",
    }
    for offset, line in enumerate(invoice.get("lines", [])):
        row = start_row + offset
        for field in col_labels:
            cell = f"{cols[field]}{row}"
            add(cell, str(line.get(field, "")), f"line {offset + 1}", col_labels[field])

    footer = layout["footer_cells"]
    add(footer["sales_tax"], str(invoice.get("sales_tax", "")), "footer")
    add(footer["invoice_total"], str(invoice.get("invoice_total", "")), "footer")
    add(footer["net_due"], str(invoice.get("net_due", invoice.get("invoice_total", ""))), "footer")
    return sequence


def copy_empty_workbook(root: Path, config: dict[str, Any], output_path: Path) -> Path:
    template = root / config["paths"]["template"]
    if not template.exists():
        create_sage_template(root, config)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, output_path)
    return output_path
